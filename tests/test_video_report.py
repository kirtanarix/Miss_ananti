"""Video export tests: invented results, isolated workbooks, no live AI calls."""

import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from copy import copy, deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import video_report as report
from score_video import RUBRIC_PATH, load_rubric, score_video


class VideoReportTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.path = self.root / "outputs" / "video_final_report.xlsx"
        root_patch = patch.object(report, "ROOT", self.root)
        root_patch.start()
        self.addCleanup(root_patch.stop)
        self.rubric = load_rubric()
        self.save(self.result("FAKE001"))

    def result(self, candidate, run=1, level=6):
        levels = dict.fromkeys(self.rubric["parameters"], level)
        marks, total = score_video(levels)
        parameters = {}
        for code in levels:
            kind = "observation" if code == "Composure" else "speech"
            parameters[code] = {
                "level": level, "reason": "A specific weekly activity supports the descriptor.",
                "reduction_basis": [], "evidence": [{"kind": kind,
                    "text": "She completes all three responses." if kind == "observation" else
                            "I practise helping neighbours every week.",
                    "start_seconds": 1.0, "end_seconds": 3.2,
                    "supports": ["task_completion"] if kind == "observation" else ["speech"]}]
            }
        return {"candidate_id": candidate, "run": run, "status": "evaluated",
                "video_rubric_version": self.rubric["version"],
                "rubric_sha256": hashlib.sha256(RUBRIC_PATH.read_bytes()).hexdigest(),
                "primary_language": "English", "overall_flags": [],
                "parameters": parameters, "review_reason": "", "review_evidence": [],
                "marks": {c: float(v) for c, v in marks.items()}, "total": float(total),
                "review_required": False, "experimental": True, "official_score": False}

    def save(self, result):
        folder = self.root / "outputs" / result["candidate_id"]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"video_ai_run{result['run']}.json"
        path.write_text(json.dumps(result), encoding="utf-8")
        return path

    def command(self, *ids):
        with redirect_stdout(io.StringIO()) as output:
            code = report.main(["--export-only", *ids])
        return code, output.getvalue()

    def workbook(self):
        wb = load_workbook(self.path)
        self.addCleanup(wb.close)
        return wb

    def snapshot(self, wb):
        return {ws.title: {
            "cells": {(c.row, c.column): (c.value, c.data_type, c.number_format,
                       copy(c.font), copy(c.fill), copy(c.border), copy(c.alignment))
                      for row in ws for c in row},
            "merges": {str(r) for r in ws.merged_cells.ranges},
            "widths": {col: dim.width for col, dim in ws.column_dimensions.items()},
            "heights": {row: dim.height for row, dim in ws.row_dimensions.items()},
            "freeze": ws.freeze_panes,
        } for ws in wb}

    def test_creates_exact_sheets_and_ai_only_headers(self):
        self.assertFalse(self.path.exists())
        self.assertEqual(self.command("FAKE001")[0], 0)
        wb = self.workbook()
        self.assertEqual(wb.sheetnames, ["Review", "Totals", "Rubric"])
        self.assertEqual([c.value for c in wb["Review"][1]], [
            "Candidate", "Parameter", "Max Marks", "AI Level", "AI Evidence", "AI Reason", "AI Marks", "Flags"])
        self.assertEqual([c.value for c in wb["Totals"][1]], [
            "Candidate", "AI SA", "AI Goals", "AI Purpose", "AI Clarity", "AI Voice", "AI Composure", "AI Total (/20)"])
        self.assertEqual(wb["Review"].max_column, 8)
        self.assertEqual(wb["Totals"].max_column, 8)
        for ws in wb:
            self.assertFalse(any(isinstance(c.value, str) and "judge" in c.value.lower()
                                 for c in ws[1]))

    def test_one_candidate_has_six_review_rows_and_one_total(self):
        self.command("FAKE001")
        wb = self.workbook()
        self.assertEqual(wb["Review"].max_row, 7)
        self.assertEqual(wb["Totals"].max_row, 2)
        self.assertEqual([r[0] for r in wb["Review"].iter_rows(min_row=2, values_only=True)], ["FAKE001"] * 6)
        self.assertEqual([r[1] for r in wb["Review"].iter_rows(min_row=2, values_only=True)], list(self.rubric["parameters"]))

    def test_levels_marks_reasons_and_total_copied(self):
        source = self.result("FAKE001")
        self.command("FAKE001")
        wb = self.workbook()
        for row, code in zip(wb["Review"].iter_rows(min_row=2), self.rubric["parameters"]):
            self.assertEqual(row[2].value, self.rubric["parameters"][code]["maximum_marks"])
            self.assertEqual(row[3].value, source["parameters"][code]["level"])
            self.assertEqual(row[5].value, source["parameters"][code]["reason"])
            self.assertEqual(row[6].value, source["marks"][code])
            self.assertEqual(row[6].number_format, "0.00")
        self.assertEqual([c.value for c in wb["Totals"][2]],
                         ["FAKE001", *source["marks"].values(), source["total"]])

    def test_rubric_copies_all_anchor_descriptors_and_reference(self):
        self.command("FAKE001")
        reference = self.workbook()["Rubric"]
        values = [c.value for row in reference for c in row]
        self.assertIn(self.rubric["version"], values)
        self.assertIn(self.rubric["source_pdf"], values)
        self.assertIn("0 / 2 / 4 / 6 / 8", values)
        for prompt in self.rubric["video_prompts"]:
            self.assertIn(prompt, values)
        for parameter in self.rubric["parameters"].values():
            for value in [parameter["name"], parameter["maximum_marks"], parameter["primary_evidence"],
                          parameter["what_earns_marks"], parameter["fairness_guardrail"], *parameter["levels"].values()]:
                self.assertIn(value, values)
        self.assertIn(self.rubric["confirmed_project_decisions"]["heavy_reading"]["rule"], values)

    def test_append_preserves_existing_cells_styles_dimensions_and_rubric(self):
        self.command("FAKE001")
        wb = self.workbook()
        before = self.snapshot(wb)
        wb.close()
        self.save(self.result("FAKE002", level=4))
        self.assertEqual(self.command("FAKE002")[0], 0)
        wb = self.workbook()
        after = self.snapshot(wb)
        for name, original in before.items():
            for cell, value in original["cells"].items():
                self.assertEqual(after[name]["cells"][cell], value)
            self.assertEqual(after[name]["widths"], original["widths"])
            self.assertEqual(after[name]["merges"], original["merges"])
            self.assertEqual(after[name]["freeze"], original["freeze"])
            for row, height in original["heights"].items():
                self.assertEqual(after[name]["heights"][row], height)
        self.assertEqual(after["Rubric"], before["Rubric"])
        self.assertEqual(wb["Review"].max_row, 13)
        self.assertEqual(wb["Totals"].max_row, 3)
        self.assertEqual(len(list((self.path.parent / "backups").glob("*.xlsx"))), 1)

    def test_duplicate_skipped_without_changing_file_or_reading_runs(self):
        self.command("FAKE001")
        original = self.path.read_bytes()
        with patch.object(report, "latest_valid") as select:
            code, output = self.command("FAKE001")
        self.assertEqual(code, 0)
        self.assertIn("already in report; skipped", output)
        select.assert_not_called()
        self.assertEqual(self.path.read_bytes(), original)

    def test_highest_numeric_valid_run_selected_without_mixing(self):
        self.save(self.result("FAKE001", run=2, level=2))
        self.save(self.result("FAKE001", run=10, level=8))
        code, output = self.command("FAKE001")
        self.assertEqual(code, 0)
        self.assertIn("video_ai_run10.json", output)
        wb = self.workbook()
        self.assertTrue(all(r[3].value == 8 for r in wb["Review"].iter_rows(min_row=2)))
        self.assertEqual(wb["Totals"]["H2"].value, 20)

    def test_latest_failed_or_corrupt_run_falls_back_to_complete_valid_run(self):
        data = self.result("FAKE001", run=2)
        data.update(status="unable_to_evaluate", parameters=None, marks=None, total=None)
        self.save(data)
        path = self.root / "outputs" / "FAKE001" / "video_ai_run3.json"
        path.write_text("broken JSON")
        code, output = self.command("FAKE001")
        self.assertEqual(code, 0)
        self.assertIn("video_ai_run1.json", output)
        self.assertEqual(self.workbook()["Totals"]["H2"].value, 15)

    def test_all_failed_statuses_create_no_report_or_scores(self):
        for status in ("unable_to_evaluate", "retry_required", "needs_review", "technical_failure", "unsupported_language"):
            with self.subTest(status=status):
                data = self.result("FAILED")
                data.update(status=status, parameters=None, marks=None, total=None)
                self.save(data)
                self.assertEqual(self.command("FAILED")[0], 1)
                self.assertFalse(self.path.exists())

    def test_failure_does_not_prevent_other_ids_appending(self):
        self.save(self.result("FAKE002"))
        self.assertEqual(self.command("FAKE001", "MISSING", "FAKE002")[0], 1)
        self.assertEqual([r[0] for r in self.workbook()["Totals"].iter_rows(min_row=2, values_only=True)], ["FAKE001", "FAKE002"])

    def test_multiple_evidence_timestamps_supports_and_reduction_preserved(self):
        data = self.result("FAKE001")
        second = {"kind": "observation", "text": "She repeatedly stops, then resumes the response.",
                  "start_seconds": 4.0, "end_seconds": 8.0, "supports": ["repeated_stopping", "recovery"]}
        item = data["parameters"]["Clarity"]
        item["evidence"].append(second)
        item["reduction_basis"] = ["repeated_stopping"]
        self.save(data)
        self.command("FAKE001")
        cell = self.workbook()["Review"]["E5"].value
        for evidence in item["evidence"]:
            self.assertIn(evidence["text"], cell)
            self.assertIn(f"{evidence['start_seconds']}–{evidence['end_seconds']}s", cell)
            for label in evidence["supports"]:
                self.assertIn(label, cell)
        self.assertIn("Reduction basis: repeated_stopping", cell)
        self.assertIn("\n\n", cell)

    def test_flags_preserved_including_evaluated_review_required(self):
        data = self.result("FAKE001")
        data.update(overall_flags=["CONTENT_SAFETY_REVIEW", "POSSIBLE_TEMPLATE_RESPONSE"], review_required=True)
        self.save(data)
        self.assertEqual(self.command("FAKE001")[0], 0)
        wb = self.workbook()
        for row in wb["Review"].iter_rows(min_row=2):
            self.assertEqual(row[7].value, "CONTENT_SAFETY_REVIEW\nPOSSIBLE_TEMPLATE_RESPONSE")
        self.assertEqual(wb["Totals"]["H2"].value, 15)

    def test_saved_marks_or_total_mismatch_rejected(self):
        for field in ("marks", "total"):
            data = self.result("BADMARKS")
            if field == "marks":
                data[field]["SA"] += 1
            else:
                data[field] += 1
            self.save(data)
            self.assertEqual(self.command("BADMARKS")[0], 1)
            self.assertFalse(self.path.exists())

    def test_identity_version_and_hash_mismatch_rejected(self):
        for key, value in (("run", 99), ("video_rubric_version", "WRONG"), ("rubric_sha256", "WRONG")):
            data = self.result("BADMETA")
            path = self.save(data)
            data[key] = value
            path.write_text(json.dumps(data))
            self.assertEqual(self.command("BADMETA")[0], 1)
            self.assertFalse(self.path.exists())
        data = self.result("BADMETA")
        path = self.save(data)
        data["candidate_id"] = "OTHER"
        path.write_text(json.dumps(data))
        self.assertEqual(self.command("BADMETA")[0], 1)

    def test_missing_parameter_invalid_level_and_missing_evidence_rejected(self):
        for defect in ("parameter", "level", "evidence"):
            data = self.result("INVALID")
            if defect == "parameter":
                del data["parameters"]["SA"]
            elif defect == "level":
                data["parameters"]["SA"]["level"] = 9
            else:
                data["parameters"]["SA"]["evidence"] = []
            self.save(data)
            self.assertEqual(self.command("INVALID")[0], 1)
            self.assertFalse(self.path.exists())

    def test_literal_formula_text_and_source_json_preserved(self):
        data = self.result("FAKE001")
        data["parameters"]["SA"]["reason"] = "=1+1"
        data["parameters"]["SA"]["evidence"][0]["text"] = "=invented speech"
        path = self.save(data)
        original = path.read_bytes()
        self.command("FAKE001")
        wb = self.workbook()
        self.assertEqual(wb["Review"]["F2"].value, "=1+1")
        self.assertEqual(wb["Review"]["F2"].data_type, "s")
        self.assertIn("=invented speech", wb["Review"]["E2"].value)
        self.assertEqual(path.read_bytes(), original)

    def test_save_or_replace_failure_keeps_existing_report(self):
        self.command("FAKE001")
        original = self.path.read_bytes()
        self.save(self.result("FAKE002"))
        for target in ("openpyxl.workbook.workbook.Workbook.save", "video_report.os.replace"):
            with patch(target, side_effect=PermissionError("test lock")):
                self.assertEqual(self.command("FAKE002")[0], 1)
            self.assertEqual(self.path.read_bytes(), original)
            self.assertEqual(list(self.path.parent.glob("tmp*.xlsx")), [])

    def test_incompatible_existing_headers_are_left_unchanged(self):
        self.command("FAKE001")
        wb = self.workbook()
        wb["Review"]["I1"] = "Unexpected column"
        wb.save(self.path)
        wb.close()
        original = self.path.read_bytes()
        self.save(self.result("FAKE002"))
        self.assertEqual(self.command("FAKE002")[0], 1)
        self.assertEqual(self.path.read_bytes(), original)

    def test_overlong_evidence_not_silently_truncated(self):
        data = self.result("TOOLONG")
        data["parameters"]["SA"]["evidence"][0]["text"] = "x" * 32768
        self.save(data)
        self.assertEqual(self.command("TOOLONG")[0], 1)
        self.assertFalse(self.path.exists())

    def test_export_never_calls_gemini_or_reads_env(self):
        with patch("score_video_ai.genai.Client") as client, patch("score_video_ai.dotenv_values") as env:
            self.assertEqual(self.command("FAKE001")[0], 0)
        client.assert_not_called()
        env.assert_not_called()

    def test_invalid_candidate_id_cannot_traverse_paths(self):
        self.assertEqual(self.command("../escape")[0], 1)
        self.assertFalse(self.path.exists())


if __name__ == "__main__":
    unittest.main()
