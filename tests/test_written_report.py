"""Final report tests: invented contestants, isolated files, mocked scoring."""

import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from copy import copy, deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import written_report as report
from score_written import load_rubric, score_written


class WrittenReportTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.path = self.root / "outputs" / "written_final_report.xlsx"
        self.rubric = load_rubric()
        self.answer = "Invented example action."
        root_patch = patch.object(report, "ROOT", self.root)
        root_patch.start()
        self.addCleanup(root_patch.stop)
        score_patch = patch.object(report, "score_contestant", side_effect=self.fake_score)
        self.scorer = score_patch.start()
        self.addCleanup(score_patch.stop)
        self.add_candidate("FAKE001")
        self.add_candidate("FAKE002")

    def add_candidate(self, candidate):
        folder = self.root / "Data" / candidate
        folder.mkdir(parents=True)
        (folder / "written_answers.json").write_text(json.dumps(dict.fromkeys(self.rubric["questions"], self.answer)))

    def result(self, candidate, run=1):
        levels = {q: {c: 4 for c, w in weights.items() if w is not None}
                  for q, weights in self.rubric["weight_matrix"].items()}
        levels["Q1"]["EV"] = 0
        calculated, total = score_written(levels, self.rubric)
        questions = {}
        for q, cells in levels.items():
            questions[q] = {"status": "valid", "levels": cells,
                            "evidence": {c: self.answer if l else "" for c, l in cells.items()},
                            "marks": float(calculated[q]["marks"]), "flags": [],
                            "answer_sha256": hashlib.sha256(self.answer.encode()).hexdigest()}
        return {"candidate_id": candidate, "run": run, "questions": questions, "written_total": float(total)}

    def fake_score(self, candidate, model):
        result = self.result(candidate)
        return result, self.saved(result)

    def saved(self, result):
        folder = self.root / "outputs" / result["candidate_id"]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"written_ai_result_run{result['run']}.json"
        path.write_text(json.dumps(result))
        return path

    def command(self, *args):
        with redirect_stdout(io.StringIO()) as output:
            code = report.main(list(args))
        return code, output.getvalue()

    def snapshot(self, wb):
        data = {}
        for ws in wb:
            data[ws.title] = {
                "cells": {(c.row, c.column): (c.value, c.data_type, c.number_format,
                                             copy(c.font), copy(c.fill), copy(c.border), copy(c.alignment))
                          for row in ws for c in row},
                "merges": {str(r) for r in ws.merged_cells.ranges},
                "widths": {col: dim.width for col, dim in ws.column_dimensions.items()},
                "heights": {row: dim.height for row, dim in ws.row_dimensions.items()},
            }
        return data

    def test_second_append_preserves_all_existing_cells_merges_widths_heights(self):
        self.assertEqual(self.command("FAKE001")[0], 0)
        wb = load_workbook(self.path)
        before = self.snapshot(wb)
        wb.close()
        self.assertEqual(self.command("FAKE002")[0], 0)
        wb = load_workbook(self.path)
        after = self.snapshot(wb)
        for name, original in before.items():
            for key, value in original["cells"].items():
                self.assertEqual(after[name]["cells"][key], value)
            self.assertTrue(original["merges"].issubset(after[name]["merges"]))
            self.assertEqual(after[name]["widths"], original["widths"])
            for key, value in original["heights"].items():
                self.assertEqual(after[name]["heights"][key], value)
        self.assertEqual(wb["Review"].max_row, 41)
        self.assertEqual(wb["Totals"].max_row, 3)
        self.assertEqual(len(wb["Review"].merged_cells.ranges), 50)
        wb.close()
        self.assertEqual(len(list((self.root / "outputs" / "backups").glob("*.xlsx"))), 1)

    def test_duplicate_refused_before_scoring(self):
        self.command("FAKE001")
        original = self.path.read_bytes()
        self.scorer.reset_mock()
        self.assertEqual(self.command("FAKE001")[0], 1)
        self.scorer.assert_not_called()
        self.assertEqual(self.path.read_bytes(), original)

    def test_invalid_input_refused_before_scoring(self):
        (self.root / "Data" / "FAKE001" / "written_answers.json").write_text('{"Q1":" "}')
        self.assertEqual(self.command("FAKE001")[0], 1)
        self.scorer.assert_not_called()
        self.assertFalse(self.path.exists())

    def test_failed_question_not_appended(self):
        result = self.result("FAKE001")
        result["questions"]["Q2"].update(status="unable_to_evaluate", marks=None, reason="API failure (ConnectError)")
        self.scorer.side_effect = None
        self.scorer.return_value = result, self.saved(result)
        code, output = self.command("FAKE001")
        self.assertEqual(code, 1)
        self.assertIn("Q2: API failure (ConnectError)", output)
        self.assertFalse(self.path.exists())
        self.assertTrue((self.root / "outputs" / "written_report_log.txt").exists())

    def test_totals_formats_and_zero_evidence_equal_calculator(self):
        self.command("FAKE001")
        result = self.result("FAKE001")
        expected, total = score_written({q: i["levels"] for q, i in result["questions"].items()}, self.rubric)
        wb = load_workbook(self.path)
        self.assertEqual([c.value for c in wb["Review"][1]], report.REVIEW_HEADERS)
        self.assertEqual([c.value for c in wb["Totals"][1]], report.TOTAL_HEADERS)
        self.assertEqual(wb["Review"].max_column, 9)
        self.assertEqual(wb["Totals"].max_column, 7)
        values = [c.value for c in wb["Totals"][2]][1:]
        self.assertEqual(values, [float(expected[q]["marks"]) for q in self.rubric["questions"]] + [float(total)])
        for row in wb["Review"].iter_rows(min_row=2):
            if row[1].value:
                q = row[1].value
                self.assertEqual(row[7].value, float(expected[q]["percent"] / 100))
                self.assertEqual(row[8].value, float(expected[q]["marks"]))
                self.assertEqual(row[7].number_format, "0.00%")
            if row[5].value == 0:
                self.assertIn(row[6].value, (None, ""))
            self.assertEqual(row[4].number_format, '0"%"')
        wb.close()

    def test_permission_error_on_save_leaves_existing_file_unchanged(self):
        self.command("FAKE001")
        original = self.path.read_bytes()
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=PermissionError("fake lock")):
            code, output = self.command("FAKE002")
        self.assertEqual(code, 1)
        self.assertIn("Close the Excel file and run: written_report.py --export-only FAKE002", output)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.glob("tmp*.xlsx")), [])

    def test_permission_error_on_replace_leaves_existing_file_unchanged(self):
        self.command("FAKE001")
        original = self.path.read_bytes()
        with patch.object(report.os, "replace", side_effect=PermissionError("fake lock")):
            self.assertEqual(self.command("FAKE002")[0], 1)
        self.assertEqual(self.path.read_bytes(), original)

    def test_marks_mismatch_not_appended(self):
        result = self.result("FAKE001")
        result["questions"]["Q1"]["marks"] += 1
        self.saved(result)
        self.assertEqual(self.command("--export-only", "FAKE001")[0], 1)
        self.scorer.assert_not_called()
        self.assertFalse(self.path.exists())

    def test_export_uses_latest_valid_run_without_scoring(self):
        self.saved(self.result("FAKE001", 2))
        invalid = self.result("FAKE001", 3)
        invalid["questions"]["Q1"]["status"] = "unable_to_evaluate"
        self.saved(invalid)
        code, output = self.command("--export-only", "FAKE001")
        self.assertEqual(code, 0)
        self.assertIn("chosen run 2", output)
        self.scorer.assert_not_called()

    def test_hash_mismatch_skips_and_legacy_accepts_with_log(self):
        result = self.result("FAKE001")
        result["questions"]["Q1"]["answer_sha256"] = "wrong"
        self.saved(result)
        self.assertEqual(self.command("--export-only", "FAKE001")[0], 1)
        self.assertFalse(self.path.exists())
        for item in result["questions"].values():
            item.pop("answer_sha256")
        self.saved(result)
        self.assertEqual(self.command("--export-only", "FAKE001")[0], 0)
        self.assertIn("accepted legacy JSON", (self.root / "outputs" / "written_report_log.txt").read_text())
        self.scorer.assert_not_called()

    def test_flags_logged_without_changing_marks(self):
        result = self.result("FAKE001")
        result["questions"]["Q1"]["flags"] = ["HARMFUL_CONTENT"]
        self.saved(result)
        self.assertEqual(self.command("--export-only", "FAKE001")[0], 0)
        contents = (self.root / "outputs" / "written_report_log.txt").read_text()
        self.assertIn("Q1 flags HARMFUL_CONTENT", contents)
        self.assertNotIn(self.answer, contents)
        wb = load_workbook(self.path)
        self.assertEqual(wb["Totals"].cell(2, 7).value, result["written_total"])
        wb.close()


if __name__ == "__main__":
    unittest.main()
