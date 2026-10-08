"""Synthetic offline tests only: no real answers, .env reads or Gemini calls."""

import copy
import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from score_written import RUBRIC_PATH, load_rubric, score_question, score_written
import score_written_ai_v2 as ai
import written_report_v2 as report

ANSWER = "I made a checklist. I listened to my teammate. We tested it and learned together."


def response(active, level=4):
    rubric = load_rubric()
    result = {field: {} for field in ai.RESPONSE_FIELDS[:-1]}
    result["flags"] = []
    for code in active:
        result["evidence"][code] = [] if level == 0 else ["I made a checklist.", "I listened to my teammate."]
        result["element_checks"][code] = [{"descriptor_element": rubric["parameters"][code]["levels"][str(level if level % 2 == 0 else level - 1)],
                                           "status": "supported", "explanation": "Synthetic structural test, not a semantic judgment."}]
        result["reasoning"][code] = "Synthetic evidence supports this test level."
        anchor = next((a for a in (0, 2, 4, 6, 8) if a > level), None)
        result["next_level_explanation"][code] = {"next_anchor_level": anchor,
            "descriptor": rubric["parameters"][code]["levels"][str(anchor)] if anchor is not None else None,
            "missing_elements": ["Synthetic missing element"] if anchor is not None else [],
            "explanation": "Synthetic next-anchor explanation." if anchor is not None else "No higher level exists."}
        result["levels"][code] = level
    return result


class ValidatorTests(unittest.TestCase):
    def validate(self, payload, answer=ANSWER):
        return ai.validate(json.dumps(payload), ["SA"], answer)

    def assertRejected(self, payload, message, answer=ANSWER):
        result = self.validate(payload, answer)
        self.assertEqual(result["status"], "unable_to_evaluate")
        self.assertIsNone(result["marks"])
        self.assertIn(message, result["reason"])

    def test_accepts_good_two_quote_response(self):
        result = self.validate(response(["SA"]))
        self.assertEqual(result["status"], "valid")
        self.assertEqual(result["evidence_verified"]["SA"], [True, True])

    def test_accepts_one_quote_and_normalizes_whitespace(self):
        item = response(["SA"])
        item["evidence"]["SA"] = ["I   made\na checklist."]
        self.assertEqual(self.validate(item)["status"], "valid")

    def test_rejects_absent_second_quote(self):
        item = response(["SA"])
        item["evidence"]["SA"][1] = "Invented synthetic statement."
        self.assertRejected(item, "quote 2 not found")

    def test_rejects_quote_over_fifteen_words(self):
        answer = " ".join(f"word{i}" for i in range(16))
        item = response(["SA"])
        item["evidence"]["SA"] = [answer]
        self.assertRejected(item, "exceeds 15 words", answer)

    def test_accepts_exactly_fifteen_words(self):
        answer = " ".join(f"word{i}" for i in range(15))
        item = response(["SA"])
        item["evidence"]["SA"] = [answer]
        self.assertEqual(self.validate(item, answer)["status"], "valid")

    def test_empty_reasoning_is_failure(self):
        for value in ("", "  ", None, 4):
            with self.subTest(value=value):
                item = response(["SA"])
                item["reasoning"]["SA"] = value
                self.assertRejected(item, "reasoning must be non-empty")

    def test_missing_next_level_explanation_is_failure(self):
        item = response(["SA"])
        del item["next_level_explanation"]
        self.assertRejected(item, "Response must contain exactly")
        item = response(["SA"])
        item["next_level_explanation"]["SA"]["explanation"] = ""
        self.assertRejected(item, "next-level explanation must be non-empty")

    def test_accepts_zero_with_empty_evidence_and_reasoning(self):
        item = response(["SA"], 0)
        self.assertEqual(self.validate(item)["status"], "valid")
        item["evidence"]["SA"] = ["I made a checklist."]
        self.assertRejected(item, "evidence must be []")

    def test_accepts_maximum_and_no_higher_anchor(self):
        item = response(["SA"], 8)
        self.assertEqual(self.validate(item)["status"], "valid")
        item["next_level_explanation"]["SA"]["next_anchor_level"] = 8
        self.assertRejected(item, "maximum level has no next anchor")

    def test_odd_level_uses_next_defined_anchor(self):
        item = response(["SA"], 5)
        self.assertEqual(self.validate(item)["status"], "valid")
        item["next_level_explanation"]["SA"]["next_anchor_level"] = 7
        self.assertRejected(item, "next_anchor_level must be 6")

    def test_unknown_flags_rejected(self):
        item = response(["SA"])
        item["flags"] = ["NEW_FLAG"]
        self.assertRejected(item, "Invalid flags")

    def test_inactive_missing_and_extra_parameters_rejected(self):
        for field in ai.RESPONSE_FIELDS[:-1]:
            item = response(["SA"])
            item[field]["CV"] = item[field]["SA"]
            self.assertRejected(item, f"{field} must contain exactly")

    def test_invalid_levels_not_defaulted(self):
        for value in (9, -1, 5.5, True, "4"):
            item = response(["SA"])
            item["levels"]["SA"] = value
            self.assertRejected(item, "level must be an integer")

    def test_element_check_required(self):
        item = response(["SA"])
        item["element_checks"]["SA"] = []
        self.assertRejected(item, "element_checks must be a non-empty list")

    def test_wrong_next_descriptor_rejected_without_semantic_judgment(self):
        item = response(["SA"])
        item["next_level_explanation"]["SA"]["descriptor"] = "Invented descriptor"
        self.assertRejected(item, "exactly match the rubric")

    def test_duplicate_keys_and_malformed_json_safe_errors(self):
        for raw in ('{"levels":{},"levels":{}}', '{"SECRET_SYNTHETIC_TEXT":'):
            result = ai.validate(raw, ["SA"], ANSWER)
            self.assertEqual(result["status"], "unable_to_evaluate")
            self.assertNotIn("SECRET_SYNTHETIC_TEXT", result["reason"])

    def test_retry_receives_validator_feedback_and_no_new_settings(self):
        bad = response(["SA"])
        bad["reasoning"]["SA"] = ""
        client = Mock()
        client.models.generate_content.side_effect = [SimpleNamespace(text=json.dumps(bad)), SimpleNamespace(text=json.dumps(response(["SA"])))]
        result = ai.evaluate_question(client, "test-model", "SYNTHETIC PROMPT", ["SA"], ANSWER, load_rubric(), sleep=lambda _: None)
        self.assertEqual(result["status"], "valid")
        self.assertEqual(result["attempts"], 2)
        sent = client.models.generate_content.call_args_list
        self.assertNotIn("VALIDATOR FEEDBACK", sent[0].kwargs["contents"])
        self.assertIn("SA: reasoning must be non-empty text", sent[1].kwargs["contents"])
        self.assertEqual(sent[1].kwargs["config"].model_dump(exclude_none=True), ai.GENERATION_CONFIG)

    def test_exhausted_retries_leave_no_marks(self):
        client = Mock()
        client.models.generate_content.return_value = SimpleNamespace(text="{}")
        result = ai.evaluate_question(client, "test", "PROMPT", ["SA"], ANSWER, load_rubric(), sleep=lambda _: None)
        self.assertEqual(result["attempts"], 3)
        self.assertEqual(result["status"], "unable_to_evaluate")
        self.assertIsNone(result["marks"])

    def test_prompt_placeholders_in_answer_not_replaced(self):
        rubric = load_rubric()
        prompt = ai.fill_prompt(ai.PROMPT.read_text(), rubric["questions"]["Q2"], "literal {question}", ["SA"], rubric)
        self.assertIn("literal {question}", prompt)
        self.assertLess(prompt.index('"reasoning"'), prompt.index('"levels"'))
        self.assertNotIn("think like a human beauty pageant judge", prompt)


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / "outputs_v2"
        self.data = self.root / "data"
        self.path = self.output / "written_final_report_v2.xlsx"
        for name, value in (("OUTPUT_ROOT", self.output),):
            patcher = patch.object(report, name, value)
            patcher.start(); self.addCleanup(patcher.stop)
        patcher = patch.object(report, "data_folder", return_value=self.data)
        patcher.start(); self.addCleanup(patcher.stop)
        self.rubric = load_rubric()

    def fixture(self, candidate="TEST", run=1):
        answers = {q: ANSWER for q in self.rubric["questions"]}
        folder = self.data / candidate
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "written_answers.json").write_text(json.dumps(answers))
        result = {"evaluation_version": "written_v2", "candidate_id": candidate, "run": run,
                  "prompt_file": "prompts/score_written_v2.txt", "rubric_file": "config/rubric_written_v1.0.yaml",
                  "rubric_version": "v1.0", "rubric_sha256": hashlib.sha256(RUBRIC_PATH.read_bytes()).hexdigest(),
                  "questions": {}, "unable_to_evaluate_count": 0, "total_is_partial": False}
        for q in self.rubric["questions"]:
            active = [c for c,w in self.rubric["weight_matrix"][q].items() if w is not None]
            item = response(active)
            item.update(status="valid", marks=float(score_question(q, item["levels"], self.rubric["weight_matrix"][q])["marks"]),
                        answer_sha256=hashlib.sha256(ai.collapse(ANSWER).encode()).hexdigest())
            result["questions"][q] = item
        result["written_total"] = float(score_written({q:x["levels"] for q,x in result["questions"].items()}, self.rubric)[1])
        folder = self.output / candidate
        folder.mkdir(parents=True, exist_ok=True)
        source = folder / f"written_ai_result_run{run}.json"
        source.write_text(json.dumps(result))
        return result, source

    def test_export_synthetic_file_adds_reasoning_and_preserves_quotes(self):
        result, source = self.fixture()
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        with patch.object(ai.genai, "Client", side_effect=AssertionError("No API")), patch.object(ai, "dotenv_values", side_effect=AssertionError("No env")), redirect_stdout(io.StringIO()):
            self.assertTrue(report.export_candidate("TEST", self.path))
        w = load_workbook(self.path)
        self.assertEqual(w.sheetnames, ["Review", "Totals", "Rubric"])
        self.assertEqual([c.value for c in w["Review"][1]], report.REVIEW_HEADERS)
        self.assertEqual(report.REVIEW_HEADERS[:9], report.V1_HEADERS)
        self.assertEqual(w["Review"].max_row, 21)
        self.assertEqual(w["Totals"].max_row, 2)
        self.assertEqual(w["Totals"]["G2"].value, 20)
        self.assertIn("Quote 1: I made a checklist.", w["Review"]["G2"].value)
        self.assertIn("Quote 2: I listened to my teammate.", w["Review"]["G2"].value)
        self.assertIn(result["questions"]["Q1"]["reasoning"]["SA"], w["Review"]["J2"].value)
        self.assertIn("Descriptor elements:", w["Review"]["J2"].value)
        self.assertIn("Next anchor 6:", w["Review"]["J2"].value)
        w.close()
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), digest)

    def test_append_and_duplicate_are_local_to_v2_workbook(self):
        self.fixture("TEST1"); self.fixture("TEST2")
        with redirect_stdout(io.StringIO()):
            report.export_candidate("TEST1", self.path)
            old = self.path.read_bytes()
            self.assertFalse(report.export_candidate("TEST1", self.path))
            self.assertEqual(self.path.read_bytes(), old)
            report.export_candidate("TEST2", self.path)
        w = load_workbook(self.path)
        self.assertEqual(w["Review"].max_row, 41)
        self.assertEqual(w["Totals"].max_row, 3)
        self.assertEqual(w["Totals"]["A2"].value, "TEST1")
        w.close()

    def test_v1_report_path_is_forbidden_and_unchanged(self):
        old = self.root / "outputs" / "written_final_report.xlsx"
        old.parent.mkdir(); old.write_bytes(b"synthetic existing v1 workbook")
        with self.assertRaisesRegex(ValueError, "v1 output paths are forbidden"):
            report.export_candidate("TEST", old)
        self.assertEqual(old.read_bytes(), b"synthetic existing v1 workbook")

    def test_highest_numeric_fully_valid_run_and_failed_fallback(self):
        self.fixture(run=2)
        result, source = self.fixture(run=10)
        answers = {q: ANSWER for q in self.rubric["questions"]}
        self.assertEqual(report.latest_valid("TEST", answers, self.rubric)[1], source)
        result["questions"]["Q2"]["reasoning"]["SA"] = ""
        source.write_text(json.dumps(result))
        self.assertEqual(report.latest_valid("TEST", answers, self.rubric)[1].name, "written_ai_result_run2.json")

    def test_failed_or_tampered_candidate_creates_no_workbook(self):
        result, source = self.fixture()
        result["questions"]["Q2"]["status"] = "unable_to_evaluate"
        source.write_text(json.dumps(result))
        with self.assertRaises(ValueError): report.export_candidate("TEST", self.path)
        self.assertFalse(self.path.exists())
        result["questions"]["Q2"]["status"] = "valid"
        result["questions"]["Q2"]["evidence"]["SA"] = ["invented"]
        source.write_text(json.dumps(result))
        with self.assertRaises(ValueError): report.export_candidate("TEST", self.path)
        self.assertFalse(self.path.exists())

    def test_provenance_and_arithmetic_mismatches_rejected(self):
        good, _ = self.fixture()
        answers = {q: ANSWER for q in self.rubric["questions"]}
        for field, value in [("rubric_sha256", "wrong"), ("written_total", 0), ("candidate_id", "OTHER")]:
            bad = copy.deepcopy(good);bad[field] = value
            with self.assertRaises(ValueError): report.prepare("TEST", bad, answers, self.rubric)

    def test_evaluator_saves_all_sent_config_and_hashes_offline(self):
        self.fixture()
        client = Mock()
        client.models.generate_content.side_effect = [SimpleNamespace(text=json.dumps(response([c for c,w in self.rubric["weight_matrix"][q].items() if w is not None]))) for q in self.rubric["questions"]]
        context = Mock()
        context.__enter__ = Mock(return_value=client)
        context.__exit__ = Mock(return_value=False)
        with patch.object(ai, "OUTPUT_ROOT", self.output), patch.object(ai, "data_folder", return_value=self.data), patch.object(ai, "dotenv_values", return_value={"GEMINI_API_KEY": "SYNTHETIC_SECRET"}), patch.object(ai.genai, "Client", return_value=context):
            result, path = ai.score_contestant("TEST", "test-model")
        self.assertEqual(result["generation_config"], ai.GENERATION_CONFIG)
        self.assertEqual(result["temperature"], 0)
        self.assertEqual(result["prompt_sha256"], hashlib.sha256(ai.PROMPT.read_bytes()).hexdigest())
        self.assertEqual(result["rubric_version"], "v1.0")
        self.assertTrue(path.is_relative_to(self.output))
        self.assertNotIn("SYNTHETIC_SECRET", path.read_text())
        self.assertEqual(result["unable_to_evaluate_count"], 0)
        for call in client.models.generate_content.call_args_list:
            self.assertEqual(call.kwargs["config"].model_dump(exclude_none=True), ai.GENERATION_CONFIG)


if __name__ == "__main__":
    unittest.main()
