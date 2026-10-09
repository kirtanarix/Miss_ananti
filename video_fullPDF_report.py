"""Export evaluated transcript-first video results into a separate AI-only workbook."""

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from score_video_fullPDF import CODES, ROOT, RUBRIC_PATH, calculate, rubric

REPORT = ROOT / "outputs" / "video_fullPDF_evaluation.xlsx"
REVIEW_HEADERS = ("Candidate", "Parameter", "Max Marks", "AI Level", "Transcript Text",
                  "Evidence Phrase", "Why This Evidence Supports the Level",
                  "Video Timestamped Behavior", "Next Anchor Gap", "AI Marks", "Candidate Flags")
TOTAL_HEADERS = ("Candidate", "AI SA", "AI Goals", "AI Purpose", "AI Clarity",
                 "AI Voice", "AI Composure", "AI Total (/20)")


def cell(sheet, row, column, value):
    target = sheet.cell(row, column)
    if isinstance(value, str):
        if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("Text cannot be preserved safely in an Excel cell")
        target.value = value
        target.data_type = "s"
    else:
        target.value = value
    target.alignment = Alignment(vertical="top", wrap_text=True)
    return target


def add_row(sheet, values):
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        cell(sheet, row, column, value)
    return row


def create_workbook(r):
    book = Workbook()
    book.active.title = "Review"
    book.create_sheet("Totals")
    book.create_sheet("Rubric")
    book.create_sheet("Audit")
    headers = {"Review": REVIEW_HEADERS, "Totals": TOTAL_HEADERS,
               "Rubric": ("Parameter", "Name", "Max Marks", "Primary Evidence",
                          "What Earns Marks", "Level", "Exact PDF Descriptor", "Fairness Guardrail"),
               "Audit": ("Candidate", "Run", "Model", "Transcript File", "Transcript SHA256",
                         "Source Video SHA256", "Rubric SHA256", "Prompt SHA256", "Decision History")}
    widths = {"Review": [15, 16, 12, 10, 65, 55, 50, 65, 29, 12, 24],
              "Totals": [15, 14, 14, 14, 14, 14, 16, 18],
              "Rubric": [15, 31, 12, 28, 48, 9, 65, 60],
              "Audit": [15, 8, 23, 58, 68, 68, 68, 68, 90]}
    from openpyxl.utils import get_column_letter
    for name, values in headers.items():
        sheet = book[name]
        add_row(sheet, values)
        sheet.freeze_panes = "A2"
        for entry in sheet[1]:
            entry.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            entry.fill = PatternFill("solid", fgColor="17365D")
        for i, width in enumerate(widths[name], 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
    for code in CODES:
        details = r["parameters"][code]
        for level in (0, 2, 4, 6, 8):
            add_row(book["Rubric"], (code, details["name"], details["maximum_marks"],
                                     details["primary_evidence"], details["what_earns_marks"],
                                     level, details["levels"][str(level)], details["fairness_guardrail"]))
    return book


def prepare(candidate):
    if not re.fullmatch(r"[A-Za-z0-9_-]+", candidate):
        raise ValueError("Invalid candidate ID")
    folder = ROOT / "outputs" / candidate
    paths = []
    for path in folder.glob("video_fullPDF_run*.json"):
        match = re.fullmatch(r"video_fullPDF_run([1-9][0-9]*)\.json", path.name)
        if match:
            paths.append((int(match.group(1)), path))
    r = rubric()
    for run, path in sorted(paths, reverse=True):
        try:
            result = json.loads(path.read_text(encoding="utf-8"))
            if result.get("status") != "evaluated" or result.get("candidate_id") != candidate or result.get("run") != run:
                continue
            if result.get("rubric_sha256") != hashlib.sha256(RUBRIC_PATH.read_bytes()).hexdigest():
                continue
            transcript = ROOT / result["transcript_path"]
            if not transcript.is_file() or hashlib.sha256(transcript.read_bytes()).hexdigest() != result["transcript_sha256"]:
                continue
            transcript_text = transcript.read_text(encoding="utf-8")
            levels = {code: result["parameters"][code]["level"] for code in CODES}
            valid_evidence = True
            for code in CODES[:5]:
                parameter = result["parameters"][code]
                if any(phrase not in transcript_text for phrase in parameter["evidence"]):
                    valid_evidence = False
                    break
                if any(item["line"] not in transcript_text or not item["explanation"].strip()
                       or not any(phrase in item["line"] for phrase in parameter["evidence"])
                       for item in parameter["transcript_evidence"]):
                    valid_evidence = False
                    break
            if not valid_evidence or any(not item.get("level_support")
                    for item in result["parameters"]["Composure"]["evidence"]):
                continue
            marks, total = calculate(levels, r)
            if any(Decimal(str(result["marks"][code])) != Decimal(str(float(marks[code]))) for code in CODES):
                continue
            if Decimal(str(result["total"])) != Decimal(str(float(total))):
                continue
            return result, path, r
        except (OSError, ValueError, TypeError, KeyError):
            continue
    raise ValueError("No complete PDF-pipeline result passes identity, transcript, rubric and arithmetic checks")


def append_candidate(book, candidate, result, r):
    review = book["Review"]
    for code in CODES:
        details = r["parameters"][code]
        parameter = result["parameters"][code]
        if code == "Composure":
            timestamped = "\n\n".join(
                f"[{item['start_seconds']:.2f}–{item['end_seconds']:.2f}s] {item['text']}"
                for item in parameter["evidence"])
            explanations = "\n\n".join(item["level_support"] for item in parameter["evidence"])
            transcript_text = evidence_phrase = ""
        else:
            transcript_text = "\n\n".join(item["line"] for item in parameter["transcript_evidence"])
            evidence_phrase = "\n".join(parameter["evidence"])
            explanations = "\n\n".join(item["explanation"] for item in parameter["transcript_evidence"])
            timestamped = ""
            if parameter.get("absence_observation"):
                explanations += "\n" + parameter["absence_observation"]
        row = add_row(review, (candidate, code, details["maximum_marks"], parameter["level"],
                               transcript_text, evidence_phrase,
                               explanations + ("\n\nOverall rationale: " + parameter["reason"] if parameter["reason"] else ""),
                               timestamped, parameter["next_anchor_gap"], result["marks"][code],
                               "\n".join(result["flags"])))
        review.cell(row, 10).number_format = "0.00"
        review.row_dimensions[row].height = max(31, 15 * math.ceil(max(
            len(transcript_text) / 64, len(evidence_phrase) / 64,
            len(explanations) / 54, len(timestamped) / 64,
            len(parameter["next_anchor_gap"]) / 49)) + 12)
    row = add_row(book["Totals"], [candidate] + [result["marks"][code] for code in CODES]
                  + [result["total"]])
    for entry in book["Totals"][row][1:]:
        entry.number_format = "0.00"
    add_row(book["Audit"], (candidate, result["run"], result["model"],
                            result["transcript_path"], result["transcript_sha256"],
                            result["source_video_sha256"], result["rubric_sha256"],
                            json.dumps(result["prompt_sha256"], ensure_ascii=False),
                            json.dumps(result["decision_history"], ensure_ascii=False)))


def export(candidate, replace_existing=False):
    result, source, r = prepare(candidate)
    book = None
    if REPORT.exists():
        current = load_workbook(REPORT)
        if current.sheetnames == ["Review", "Totals", "Rubric", "Audit"] and \
                [cell.value for cell in current["Review"][1]] == list(REVIEW_HEADERS):
            book = current
        else:
            current.close()
            backups = REPORT.parent / "backups"
            backups.mkdir(parents=True, exist_ok=True)
            backup = backups / "video_fullPDF_evaluation_before_detail_columns.xlsx"
            if not backup.exists():
                shutil.copy2(REPORT, backup)
            book = create_workbook(r)
            # Rebuild the separate workbook from each candidate's newest complete detailed run.
            for candidate_dir in sorted((ROOT / "outputs").iterdir()):
                if not candidate_dir.is_dir():
                    continue
                if not any(candidate_dir.glob("video_fullPDF_run*.json")):
                    continue
                try:
                    prior, _, prior_rubric = prepare(candidate_dir.name)
                    append_candidate(book, candidate_dir.name, prior, prior_rubric)
                except (OSError, ValueError, TypeError, KeyError):
                    continue
    else:
        book = create_workbook(r)
    try:
        if book.sheetnames != ["Review", "Totals", "Rubric", "Audit"]:
            raise ValueError("Existing workbook has unexpected sheets")
        present = any(row[0] == candidate for row in book["Totals"].iter_rows(min_row=2, max_col=1, values_only=True))
        if present and not replace_existing:
            print(f"{candidate}: already present; skipped")
            return False
        if present:
            backups = REPORT.parent / "backups"
            backups.mkdir(parents=True, exist_ok=True)
            backup = backups / f"video_fullPDF_evaluation_before_{candidate}_refresh_run{result['run']}.xlsx"
            if not backup.exists():
                shutil.copy2(REPORT, backup)
            for sheet_name in ("Review", "Totals", "Audit"):
                sheet = book[sheet_name]
                for row_index in range(sheet.max_row, 1, -1):
                    if sheet.cell(row_index, 1).value == candidate:
                        sheet.delete_rows(row_index)
        append_candidate(book, candidate, result, r)
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=REPORT.parent, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        try:
            book.save(temporary)
            os.replace(temporary, REPORT)
        finally:
            if temporary.exists():
                temporary.unlink()
        print(f"{candidate}: exported {source.name}; total {result['total']:.2f} / 20")
        return True
    finally:
        book.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--export-only", required=True, metavar="ID")
    parser.add_argument("--replace-existing", action="store_true",
                        help="Replace this candidate's rows with the newest complete detailed run")
    args = parser.parse_args(argv)
    try:
        export(args.export_only, replace_existing=args.replace_existing)
    except (OSError, ValueError, TypeError, KeyError) as error:
        parser.exit(1, f"Export failed: {type(error).__name__}: {error}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
