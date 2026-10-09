"""Export saved V3 results to the consolidated AI-only video workbook; no API calls.

Usage: .venv\\Scripts\\python.exe video_strictPDF_report.py --export-only C001 C002
Select the highest numeric run passing evaluated-status, V3-contract, identity,
rubric and deterministic-mark checks. Never mix runs. Skip already reported IDs.
Evaluated results with review flags retain their scores and flags. Unassessable
results are omitted and reported in the console, never replaced with zero marks.
"""

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from score_video_strictPDF_ai import RUBRIC_PATH, load_rubric
from score_video_strictPDF_ai import VideoResponse, unique_object, validate_response

ROOT = Path(__file__).resolve().parent
REVIEW_HEADERS = ["Candidate", "Parameter", "Max Marks", "AI Level", "AI Evidence",
                  "AI Reason", "AI Marks", "Flags"]
TOTAL_HEADERS = ["Candidate", "AI SA", "AI Goals", "AI Purpose", "AI Clarity",
                 "AI Voice", "AI Composure", "AI Total (/20)"]


def write_cell(cell, value):
    if value == "":
        cell.value = None  # Stable blank cells across save/load/append cycles.
        return
    if isinstance(value, str):
        if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("Source text cannot be preserved in an Excel cell")
        cell.value = value
        cell.data_type = "s"  # Evidence/reasons starting with '=' stay literal text.
    else:
        cell.value = value


def style_row(sheet, row, header=False):
    side = Side(style="thin", color="B8C4D0")
    for cell in sheet[row]:
        cell.font = Font(name="Arial", size=10, bold=header,
                         color="FFFFFF" if header else "000000")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        if header:
            cell.fill = PatternFill("solid", fgColor="17365D")


def reference_rows(rubric):
    """Readable reference blocks, copying descriptors and operational wording."""
    yield ["Video rubric", rubric["version"]]
    yield ["Source", rubric["source_pdf"]]
    yield ["Level scale", "0 / 2 / 4 / 6 / 8"]
    yield ["Odd levels", rubric["confirmed_project_decisions"]["odd_levels"]["rule"]]
    yield ["Video prompts", "Prompt text"]
    for index, prompt in enumerate(rubric["video_prompts"], 1):
        yield [index, prompt]
    yield ["Context", rubric["context_guardrail"]]
    yield []
    yield ["Parameter", "Name", "Max Marks", "Primary evidence", "What earns marks", "Fairness guardrail"]
    for code, parameter in rubric["parameters"].items():
        yield [code, parameter["name"], parameter["maximum_marks"], parameter["primary_evidence"],
               parameter["what_earns_marks"], parameter["fairness_guardrail"]]
    yield []
    yield ["Parameter", "0", "2", "4", "6", "8"]
    for code, parameter in rubric["parameters"].items():
        yield [code, *[parameter["levels"][str(level)] for level in (0, 2, 4, 6, 8)]]
    yield []
    yield rubric["prompt_to_parameter_evidence_map"]["columns"]
    yield from rubric["prompt_to_parameter_evidence_map"]["rows"]
    yield ["Evidence map guidance", rubric["prompt_to_parameter_evidence_map"]["guidance"]]
    yield []
    yield ["PDF prohibited factors", "Strict PDF delivery scope: content-only scoring; task delivery only for Composure."]
    for factor in rubric["prohibited_judging_factors"]["factors"]:
        yield ["PDF wording", factor]
    yield []
    yield ["Confirmed project decisions", "Current operational rule/value"]
    for name, decision in rubric["confirmed_project_decisions"].items():
        if isinstance(decision, dict):
            if "rule" in decision:
                yield [name, decision["rule"]]
            for key, value in decision.items():
                if key in {"rule", "affected_pdf_rules"}:
                    continue
                if isinstance(value, list):
                    value = "\n".join(map(str, value))
                elif value is None:
                    value = "No current limit" if name == "duration" else "Not specified"
                yield [f"{name}: {key}", value]
            for value in decision.get("affected_pdf_rules", []):
                yield ["Confirmed precedence", value]
        else:
            yield [name, decision]
    yield []
    yield ["Flag", "Existing meaning"]
    for group in ("non_scoring_flags", "scoring_related_flags"):
        for flag, description in rubric[group].items():
            yield [flag, description]
    yield []
    yield ["Unresolved items", "No additional rules supplied"]
    for item in rubric["ambiguities_for_bd"]:
        yield ["Existing TODO", item]


def create_workbook(rubric):
    wb = Workbook()
    wb.active.title = "Review"
    wb.create_sheet("Totals")
    wb.create_sheet("Rubric")
    for name, headers, widths in (
            ("Review", REVIEW_HEADERS, [15, 17, 12, 12, 68, 52, 12, 32]),
            ("Totals", TOTAL_HEADERS, [15, 14, 14, 14, 14, 14, 16, 19])):
        sheet = wb[name]
        sheet.append(headers)
        style_row(sheet, 1, True)
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "A2"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    reference = wb["Rubric"]
    for values in reference_rows(rubric):
        row = reference.max_row + 1 if reference["A1"].value is not None else 1
        if not values:
            reference.cell(row, 1, None)
            continue
        for column, value in enumerate(values, 1):
            write_cell(reference.cell(row, column), value)
        is_header = values[0] in {"Video prompts", "Parameter", "Prompt", "PDF prohibited factors",
                                  "Confirmed project decisions", "Flag", "Unresolved items"}
        style_row(reference, row, is_header)
        reference.row_dimensions[row].height = max(28, 14 * max(
            sum(max(1, math.ceil(len(line) / 38)) for line in str(value).splitlines())
            for value in values) + 10)
    for index in range(1, reference.max_column + 1):
        reference.column_dimensions[get_column_letter(index)].width = 42 if index > 1 else 30
    reference.freeze_panes = "A2"
    return wb


def check_workbook(wb):
    if wb.sheetnames != ["Review", "Totals", "Rubric"]:
        raise ValueError("Report must contain exactly Review, Totals, Rubric")
    for name, headers in (("Review", REVIEW_HEADERS), ("Totals", TOTAL_HEADERS)):
        if wb[name].max_column != len(headers) or [cell.value for cell in wb[name][1]] != headers:
            raise ValueError(f"Unexpected {name} columns; report left unchanged")


def has_candidate(wb, candidate):
    return any(row[0] == candidate for name in ("Review", "Totals")
               for row in wb[name].iter_rows(min_row=2, max_col=1, values_only=True))


def prepare(candidate, result, run, rubric):
    if not isinstance(result, dict) or result.get("status") != "evaluated":
        raise ValueError("Run is not an evaluated result")
    if result.get("candidate_id") != candidate or type(result.get("run")) is not int or result["run"] != run:
        raise ValueError("Candidate ID or filename/run mismatch")
    if result.get("video_rubric_version") != rubric["version"]:
        raise ValueError("Video rubric version mismatch")
    current_hash = hashlib.sha256(RUBRIC_PATH.read_bytes()).hexdigest()
    if result.get("rubric_sha256") and result["rubric_sha256"] != current_hash:
        raise ValueError("Video rubric hash mismatch")
    # Reuse the existing pure V3 contract validator and its V1 calculator call.
    ai_fields = {key: result[key] for key in VideoResponse.model_fields if key in result}
    validated = validate_response(json.dumps(ai_fields, ensure_ascii=False))
    if validated["status"] != "evaluated" or validated["total"] is None:
        raise ValueError("Run does not pass the existing V3 output contract")
    saved_marks = result.get("marks")
    if not isinstance(saved_marks, dict) or set(saved_marks) != set(rubric["parameters"]):
        raise ValueError("Missing or invalid saved marks")
    for code, expected in validated["marks"].items():
        value = saved_marks[code]
        if type(value) not in (int, float) or not math.isfinite(value) or Decimal(str(value)) != Decimal(str(expected)):
            raise ValueError(f"{code}: saved marks differ from V1 calculator")
    total = result.get("total")
    if type(total) not in (int, float) or not math.isfinite(total) or Decimal(str(total)) != Decimal(str(validated["total"])):
        raise ValueError("Saved total differs from V1 calculator")
    return result


def latest_valid(candidate, rubric):
    paths = []
    for path in (ROOT / "outputs" / candidate).glob("video_strictPDF_ai_run*.json"):
        match = re.fullmatch(r"video_strictPDF_ai_run([1-9][0-9]*)\.json", path.name)
        if match:
            paths.append((int(match.group(1)), path))
    for run, path in sorted(paths, reverse=True):
        try:
            result = json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=unique_object)
            return prepare(candidate, result, run, rubric), path
        except (OSError, ValueError, TypeError, KeyError):
            continue
    raise ValueError("No saved evaluated run passes V3, identity, rubric and V1 marks checks; not added to Excel")


def evidence_text(parameter):
    blocks = []
    for item in parameter["evidence"]:
        start, end = str(item["start_seconds"]), str(item["end_seconds"])
        blocks.append(f"[{item['kind'].capitalize()} {start}–{end}s]\n{item['text']}\n"
                      + "Supports: " + ", ".join(item["supports"]))
    if parameter["reduction_basis"]:
        blocks.append("Reduction basis: " + ", ".join(parameter["reduction_basis"]))
    return "\n\n".join(blocks)


def append_rows(wb, candidate, result, rubric):
    review = wb["Review"]
    for code, definition in rubric["parameters"].items():
        parameter = result["parameters"][code]
        values = [candidate, code, definition["maximum_marks"], parameter["level"],
                  evidence_text(parameter), parameter["reason"], result["marks"][code],
                  "\n".join(result["overall_flags"])]
        row = review.max_row + 1
        for column, value in enumerate(values, 1):
            write_cell(review.cell(row, column), value)
        style_row(review, row)
        review.cell(row, 3).number_format = "0.00"
        review.cell(row, 4).number_format = "0"
        review.cell(row, 7).number_format = "0.00"
        lines = max(sum(max(1, math.ceil(len(line) / width)) for line in value.splitlines()) or 1
                    for value, width in ((values[4], 65), (values[5], 49), (values[7], 29)))
        review.row_dimensions[row].height = max(30, 15 * lines + 12)
    totals = wb["Totals"]
    totals.append([candidate, *[result["marks"][code] for code in rubric["parameters"]], result["total"]])
    style_row(totals, totals.max_row)
    totals.row_dimensions[totals.max_row].height = 28
    for cell in list(totals[totals.max_row])[1:]:
        cell.number_format = "0.00"


def save_safely(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        backups = path.parent / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        now = datetime.now(timezone(timedelta(hours=5, minutes=30)))
        shutil.copy2(path, backups / f"{path.stem}_{now.strftime('%Y%m%d_%H%M%S_%f')}.xlsx")
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def export_candidate(candidate, path):
    if not re.fullmatch(r"[A-Za-z0-9_-]+", candidate):
        raise ValueError("Invalid candidate ID")
    rubric = load_rubric()
    wb = load_workbook(path) if path.exists() else create_workbook(rubric)
    try:
        check_workbook(wb)
        if has_candidate(wb, candidate):
            print(f"{candidate}: already in report; skipped")
            return False
        result, source = latest_valid(candidate, rubric)
        append_rows(wb, candidate, result, rubric)
        save_safely(wb, path)
        print(f"{candidate}: exported {source.name}; AI total {result['total']:.2f} / 20")
        return True
    finally:
        wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--export-only", nargs="+", required=True, metavar="ID")
    args = parser.parse_args(argv)
    path = ROOT / "outputs" / "video_strictPDF_evaluation.xlsx"
    failed = False
    for candidate in args.export_only:
        try:
            export_candidate(candidate, path)
        except PermissionError:
            print(f"{candidate}: cannot save; close Excel and rerun video_strictPDF_report.py --export-only {candidate}")
            failed = True
        except (OSError, ValueError, KeyError, TypeError):
            print(f"{candidate}: no compatible evaluated result or report could not be read/preserved; not added")
            failed = True
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
