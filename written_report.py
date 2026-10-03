"""Append AI-only written evaluations to the final workbook."""

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from score_written import load_rubric, score_written
from score_written_ai import collapse, score_contestant

DEFAULT_MODEL = "gemini-3.8-flash"
ROOT = Path(__file__).resolve().parent
REVIEW_HEADERS = ["Candidate", "Question", "Answer (candidate's text)", "Parameter",
                  "Weight %", "AI level", "Evidence", "AI question score %", "AI marks (out of 8)"]
TOTAL_HEADERS = ["Candidate", "AI Q1", "AI Q2", "AI Q3", "AI Q4", "AI Q5", "AI total (/40)"]


def now():
    return datetime.now(timezone(timedelta(hours=5, minutes=30)))


def log(candidate, message):
    folder = ROOT / "outputs"
    folder.mkdir(parents=True, exist_ok=True)
    with (folder / "written_report_log.txt").open("a", encoding="utf-8") as stream:
        stream.write(f"{now().isoformat()} {candidate} {message}\n")


def answers_for(candidate, rubric):
    if not candidate or Path(candidate).name != candidate or candidate in {".", ".."}:
        raise ValueError("Invalid contestant ID")
    folder = next((ROOT / name for name in ("Data", "data") if (ROOT / name).is_dir()), None)
    if folder is None:
        raise ValueError("No Data or data folder")
    try:
        answers = json.loads((folder / candidate / "written_answers.json").read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        raise ValueError("Missing or invalid written_answers.json") from None
    if not isinstance(answers, dict) or any(not isinstance(answers.get(q), str) or not answers[q].strip()
                                           for q in rubric["questions"]):
        raise ValueError("Q1 to Q5 must each be a non-empty string")
    return answers


def text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Submitted text beginning with '=' must never become a formula.


def style_row(sheet, row, header=False):
    side = Side(style="thin", color="B8C4D0")
    for cell in sheet[row]:
        cell.font = Font(name="Arial", size=10, bold=header, color="FFFFFF" if header else "000000")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        if header:
            cell.fill = PatternFill("solid", fgColor="17365D")


def create_workbook(rubric):
    wb = Workbook()
    review = wb.active
    review.title = "Review"
    totals = wb.create_sheet("Totals")
    reference = wb.create_sheet("Rubric")
    for sheet, headers, widths in ((review, REVIEW_HEADERS, [15, 12, 72, 16, 12, 12, 48, 20, 20]),
                                   (totals, TOTAL_HEADERS, [15, 15, 15, 15, 15, 15, 20])):
        sheet.append(headers)
        style_row(sheet, 1, True)
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 32
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    reference.append(["Question", "Question text"])
    for q, question in rubric["questions"].items():
        reference.append([q, question])
    reference.append([])
    weight_header = reference.max_row + 1
    reference.append(["Question", *rubric["parameters"]])
    for q, weights in rubric["weight_matrix"].items():
        reference.append([q, *[int(weights[c].removesuffix("%")) if weights[c] is not None else None
                               for c in rubric["parameters"]]])
        for cell in list(reference[reference.max_row])[1:6]:
            cell.number_format = '0"%"'
    reference.append([])
    level_header = reference.max_row + 1
    reference.append(["Parameter", "Name", "what_earns_marks", "0", "2", "4", "6", "8"])
    for code, parameter in rubric["parameters"].items():
        reference.append([code, parameter["name"], parameter["what_earns_marks"],
                          *[parameter["levels"][str(level)] for level in (0, 2, 4, 6, 8)]])
    reference.append([])
    exclusion_header = reference.max_row + 1
    reference.append(["not_scored"])
    for item in rubric["not_scored"]:
        row = reference.max_row + 1
        reference.cell(row, 1, item)
        reference.merge_cells(start_row=row, end_row=row, start_column=1, end_column=8)
    for row in range(1, reference.max_row + 1):
        style_row(reference, row, row in {1, weight_header, level_header, exclusion_header})
        reference.row_dimensions[row].height = 90 if row > level_header else 65
    for col in range(1, 9):
        reference.column_dimensions[get_column_letter(col)].width = 40 if col > 1 else 18
    reference.freeze_panes = "A2"
    return wb


def check_workbook(wb):
    if wb.sheetnames != ["Review", "Totals", "Rubric"]:
        raise ValueError("Workbook must contain Review, Totals, Rubric")
    for name, headers in (("Review", REVIEW_HEADERS), ("Totals", TOTAL_HEADERS)):
        if wb[name].max_column != len(headers) or [c.value for c in wb[name][1]] != headers:
            raise ValueError(f"Unexpected {name} columns")


def has_candidate(wb, candidate):
    return any(row[0] == candidate for name in ("Review", "Totals")
               for row in wb[name].iter_rows(min_row=2, max_col=1, values_only=True))


def latest_valid(candidate, rubric):
    paths = list((ROOT / "outputs" / candidate).glob("written_ai_result_run*.json"))
    paths.sort(key=lambda p: int(re.search(r"run(\d+)$", p.stem).group(1)), reverse=True)
    for path in paths:
        try:
            result = json.loads(path.read_text(encoding="utf-8-sig"))
            questions = result["questions"]
            if all(questions[q]["status"] == "valid" for q in rubric["questions"]):
                return result, path
        except (OSError, ValueError, KeyError, TypeError):
            continue
    raise ValueError("No saved run with five valid questions")


def prepare(candidate, result, answers, rubric):
    if result.get("candidate_id") != candidate:
        raise ValueError("Saved contestant ID mismatch")
    questions = result.get("questions", {})
    failures = []
    for q in rubric["questions"]:
        item = questions.get(q, {})
        if item.get("status") != "valid":
            # Print the scorer's saved diagnostics, not any candidate text.
            reason = item.get("reason", "Missing or invalid question result")
            failures.append(f"{q}: {reason}")
    if failures:
        raise ValueError("; ".join(failures))
    missing_hash = False
    for q in rubric["questions"]:
        item = questions[q]
        expected = hashlib.sha256(collapse(answers[q]).encode("utf-8")).hexdigest()
        if "answer_sha256" in item:
            if item["answer_sha256"] != expected:
                raise ValueError(f"{q}: answer_sha256 mismatch")
        else:
            missing_hash = True
        active = {c for c, w in rubric["weight_matrix"][q].items() if w is not None}
        if set(item["levels"]) != active or any(type(v) is not int for v in item["levels"].values()):
            raise ValueError(f"{q}: invalid active levels")
        for c, level in item["levels"].items():
            phrase = item["evidence"].get(c)
            if not isinstance(phrase, str) or (level == 0 and phrase != ""):
                raise ValueError(f"{q} {c}: invalid evidence")
            # Legacy JSONs without hashes are accepted as requested.
    calculated, total = score_written({q: questions[q]["levels"] for q in rubric["questions"]}, rubric)
    for q, calculated_question in calculated.items():
        if Decimal(str(questions[q].get("marks"))) != calculated_question["marks"]:
            raise ValueError(f"{q}: saved marks differ from calculator")
    return calculated, total, missing_hash


def append_rows(wb, candidate, answers, result, rubric, calculated, total):
    review = wb["Review"]
    for q in rubric["questions"]:
        start = review.max_row + 1
        active = [(c, w) for c, w in rubric["weight_matrix"][q].items() if w is not None]
        for c, weight in active:
            row = review.max_row + 1
            values = [candidate, q, answers[q], c, int(weight.removesuffix("%")),
                      result["questions"][q]["levels"][c], result["questions"][q]["evidence"][c],
                      float(calculated[q]["percent"] / 100), float(calculated[q]["marks"])]
            for col, value in enumerate(values, 1):
                text(review.cell(row, col), value)
            style_row(review, row)
            review.cell(row, 5).number_format = '0"%"'
            review.cell(row, 8).number_format = "0.00%"
            review.cell(row, 9).number_format = "0.00"
            evidence_lines = sum(max(1, math.ceil(len(line) / 45)) for line in values[6].splitlines()) or 1
            answer_lines = sum(max(1, math.ceil(len(line) / 68)) for line in answers[q].splitlines()) or 1
            review.row_dimensions[row].height = max(30, 14 * evidence_lines + 8,
                                                    (14 * answer_lines + 12) / len(active))
        for col in (1, 2, 3, 8, 9):
            review.merge_cells(start_row=start, end_row=review.max_row, start_column=col, end_column=col)
    totals = wb["Totals"]
    totals.append([candidate, *[float(calculated[q]["marks"]) for q in rubric["questions"]], float(total)])
    style_row(totals, totals.max_row)
    totals.row_dimensions[totals.max_row].height = 25
    for cell in list(totals[totals.max_row])[1:]:
        cell.number_format = "0.00"


def save_safely(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        backups = ROOT / "outputs" / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        backup = backups / f"{path.stem}_{now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
        shutil.copy2(path, backup)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def c005_replacement_start(wb):
    """Allow the explicitly requested correction only for the final C005 block."""
    review, totals = wb["Review"], wb["Totals"]
    starts = [row[0].row for row in review if row[0].value == "C005"]
    total_rows = [row[0].row for row in totals if row[0].value == "C005"]
    if (len(starts) != 5 or starts != [starts[0] + n for n in (0, 5, 9, 13, 17)]
            or starts[0] + 19 != review.max_row or total_rows != [totals.max_row]):
        raise ValueError("C005 replacement requires exactly one final 20-row block and one final Totals row")
    return starts[0]


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("candidate", nargs="?")
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--export-only", nargs="+", metavar="ID")
    parser.add_argument("--replace-c005", action="store_true",
                        help="Explicitly replace the existing final C005 block after corrected-answer scoring")
    args = parser.parse_args(argv)
    if args.replace_c005 and not (args.candidate == "C005" or args.export_only == ["C005"]):
        parser.error("--replace-c005 is allowed only for C005")
    if bool(args.candidate) == bool(args.export_only):
        parser.error("Supply one contestant ID or --export-only IDs")
    path = args.report or ROOT / "outputs" / "written_final_report.xlsx"
    rubric = load_rubric()
    try:
        wb = load_workbook(path) if path.exists() else create_workbook(rubric)
        check_workbook(wb)
    except (OSError, ValueError):
        print("Cannot read the report workbook")
        return 1
    failed = False
    try:
        for candidate in args.export_only or [args.candidate]:
            try:
                answers = answers_for(candidate, rubric)
                replacing = args.replace_c005 and has_candidate(wb, candidate)
                replacement_start = c005_replacement_start(wb) if replacing else None
                if has_candidate(wb, candidate) and not replacing:
                    print(f"{candidate}: already in report; skipped")
                    if not args.export_only:
                        failed = True
                    continue
                result, source = (latest_valid(candidate, rubric) if args.export_only
                                  else score_contestant(candidate, args.model))
                if args.export_only:
                    print(f"{candidate}: chosen run {result['run']}")
                # Re-read after scoring so an input change is caught by the fingerprint.
                answers = answers_for(candidate, rubric)
                calculated, total, missing_hash = prepare(candidate, result, answers, rubric)
                if missing_hash:
                    log(candidate, f"run {result['run']}: answer_sha256 absent; accepted legacy JSON")
                for q, item in result["questions"].items():
                    if item.get("flags"):
                        flags = ", ".join(item["flags"])
                        print(f"{candidate} {q}: warning flags {flags}")
                        log(candidate, f"{q} flags {flags}")
                if replacing:
                    # Remove only the verified final C005 block in memory; disk stays
                    # unchanged until the existing backup and atomic-save path succeeds.
                    for merged in list(wb["Review"].merged_cells.ranges):
                        if merged.min_row >= replacement_start:
                            wb["Review"].unmerge_cells(str(merged))
                    wb["Review"].delete_rows(replacement_start, 20)
                    wb["Totals"].delete_rows(wb["Totals"].max_row, 1)
                append_rows(wb, candidate, answers, result, rubric, calculated, total)
                try:
                    save_safely(wb, path)
                except PermissionError:
                    suffix = " --replace-c005" if replacing else ""
                    print(f"Close the Excel file and run: written_report.py --export-only {candidate}{suffix}")
                    return 1
                print(f"{candidate}: run {result['run']}, AI total {total:.2f}")
            except (ValueError, OSError, KeyError, TypeError, ArithmeticError) as error:
                # These errors are generated locally; never display SDK exceptions.
                message = str(error) if isinstance(error, ValueError) else type(error).__name__
                print(f"{candidate}: {message}")
                log(candidate, message)
                failed = True
                # Discard any unsaved in-memory append before processing the next ID.
                wb.close()
                wb = load_workbook(path) if path.exists() else create_workbook(rubric)
        return 1 if failed else 0
    finally:
        wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
