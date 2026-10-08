"""Export saved written v2 JSON only; never evaluate or touch a v1 workbook.

Run: python written_report_v2.py --export-only ID
Default: outputs_v2/written_final_report_v2.xlsx
"""

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from score_written import RUBRIC_PATH, load_rubric, score_written
from score_written_ai_v2 import OUTPUT_ROOT, RESPONSE_FIELDS, collapse, data_folder, read_answers, unique_object, validate
from written_report import REVIEW_HEADERS as V1_HEADERS, TOTAL_HEADERS, create_workbook as create_v1_workbook, has_candidate, style_row, text

ROOT = Path(__file__).resolve().parent
REPORT_PATH = OUTPUT_ROOT / "written_final_report_v2.xlsx"
REVIEW_HEADERS = [*V1_HEADERS, "AI reasoning"]


def create_workbook(rubric):
    wb = create_v1_workbook(rubric)
    review = wb["Review"]
    text(review.cell(1, 10), "AI reasoning")
    style_row(review, 1, True)
    review.column_dimensions["J"].width = 85
    return wb


def check_workbook(wb):
    if wb.sheetnames != ["Review", "Totals", "Rubric"]:
        raise ValueError("V2 report requires exactly Review, Totals, Rubric")
    for name, headers in (("Review", REVIEW_HEADERS), ("Totals", TOTAL_HEADERS)):
        if wb[name].max_column != len(headers) or [c.value for c in wb[name][1]] != headers:
            raise ValueError(f"Unexpected {name} columns; not a v2 report")


def prepare(candidate, result, answers, rubric):
    if (result.get("evaluation_version") != "written_v2" or result.get("candidate_id") != candidate
            or result.get("prompt_file") != "prompts/score_written_v2.txt"
            or result.get("rubric_file") != "config/rubric_written_v1.0.yaml"
            or result.get("rubric_version") != "v1.0"
            or result.get("rubric_sha256") != hashlib.sha256(RUBRIC_PATH.read_bytes()).hexdigest()):
        raise ValueError("V2 identity or rubric provenance mismatch")
    questions = result.get("questions")
    if not isinstance(questions, dict) or set(questions) != set(rubric["questions"]):
        raise ValueError("Five v2 question results are required")
    for q in rubric["questions"]:
        item = questions[q]
        if not isinstance(answers.get(q), str) or not answers[q].strip():
            raise ValueError(f"{q}: non-empty candidate answer required")
        if item.get("status") != "valid":
            raise ValueError(f"{q}: failed evaluation; no partial candidate exported")
        if item.get("answer_sha256") != hashlib.sha256(collapse(answers[q]).encode()).hexdigest():
            raise ValueError(f"{q}: answer fingerprint mismatch")
        active = [c for c, weight in rubric["weight_matrix"][q].items() if weight is not None]
        raw = {field: item[field] for field in RESPONSE_FIELDS if field in item}
        checked = validate(json.dumps(raw, ensure_ascii=False), active, answers[q], rubric)
        if checked["status"] != "valid":
            raise ValueError(f"{q}: saved v2 response fails validation: " + checked["reason"])
    calculated, total = score_written({q: questions[q]["levels"] for q in rubric["questions"]}, rubric)
    for q in rubric["questions"]:
        saved = questions[q].get("marks")
        if type(saved) not in (int, float) or Decimal(str(saved)) != calculated[q]["marks"]:
            raise ValueError(f"{q}: saved marks mismatch")
    if result.get("total_is_partial") is not False or result.get("unable_to_evaluate_count") != 0:
        raise ValueError("Partial result cannot be exported")
    saved_total = result.get("written_total")
    if type(saved_total) not in (int, float) or Decimal(str(saved_total)) != total:
        raise ValueError("Saved total mismatch")
    return calculated, total


def latest_valid(candidate, answers, rubric):
    paths = []
    for path in (OUTPUT_ROOT / candidate).glob("written_ai_result_run*.json"):
        match = re.fullmatch(r"written_ai_result_run([1-9][0-9]*)\.json", path.name)
        if match:
            paths.append((int(match[1]), path))
    for run, path in sorted(paths, reverse=True):
        try:
            result = json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=unique_object)
            if type(result.get("run")) is not int or result["run"] != run:
                continue
            calculated, total = prepare(candidate, result, answers, rubric)
            return result, path, calculated, total
        except (OSError, ValueError, KeyError, TypeError, ArithmeticError):
            continue
    raise ValueError("No saved v2 run passes all five questions, provenance, evidence and marks checks")


def reasoning_text(item, code):
    blocks = ["Descriptor elements:"]
    for check in item["element_checks"][code]:
        blocks.append(f"{check['descriptor_element']} [{check['status']}]\n{check['explanation']}")
    blocks.append("AI reasoning:\n" + item["reasoning"][code])
    nxt = item["next_level_explanation"][code]
    if nxt["next_anchor_level"] is None:
        blocks.append("Next level:\n" + nxt["explanation"])
    else:
        blocks.append(f"Next anchor {nxt['next_anchor_level']}:\n{nxt['descriptor']}\n"
                      + "Missing elements: " + "; ".join(nxt["missing_elements"]) + "\n" + nxt["explanation"])
    return "\n\n".join(blocks)


def append_rows(wb, candidate, answers, result, rubric, calculated, total):
    review = wb["Review"]
    for q in rubric["questions"]:
        start = review.max_row + 1
        item = result["questions"][q]
        active = [(c, weight) for c, weight in rubric["weight_matrix"][q].items() if weight is not None]
        for code, weight in active:
            values = [candidate, q, answers[q], code, int(weight.removesuffix("%")), item["levels"][code],
                      "\n\n".join(f"Quote {index}: {quote}" for index, quote in enumerate(item["evidence"][code], 1)),
                      float(calculated[q]["percent"] / 100), float(calculated[q]["marks"]), reasoning_text(item, code)]
            row = review.max_row + 1
            for col, value in enumerate(values, 1):
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError("Source text exceeds Excel cell limit; not truncated")
                text(review.cell(row, col), value)
            style_row(review, row)
            review.cell(row, 5).number_format = '0"%"'
            review.cell(row, 8).number_format = "0.00%"
            review.cell(row, 9).number_format = "0.00"
            lines = max(sum(max(1, math.ceil(len(line) / width)) for line in value.splitlines()) or 1
                        for value, width in ((values[6], 45), (values[9], 80)))
            review.row_dimensions[row].height = max(30, 14 * lines + 12)
        for col in (1, 2, 3, 8, 9):
            review.merge_cells(start_row=start, end_row=review.max_row, start_column=col, end_column=col)
    totals = wb["Totals"]
    totals.append([candidate, *[float(calculated[q]["marks"]) for q in rubric["questions"]], float(total)])
    style_row(totals, totals.max_row)
    for cell in list(totals[totals.max_row])[1:]:
        cell.number_format = "0.00"


def safe_report_path(path):
    path = Path(path).resolve()
    if not path.is_relative_to(OUTPUT_ROOT.resolve()) or path.suffix.lower() != ".xlsx":
        raise ValueError("V2 workbook must be an .xlsx under outputs_v2; v1 output paths are forbidden")
    return path


def save_safely(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        backups = OUTPUT_ROOT / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y%m%d_%H%M%S_%f")
        shutil.copy2(path, backups / f"{path.stem}_{stamp}.xlsx")
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def export_candidate(candidate, path=REPORT_PATH):
    if not re.fullmatch(r"[A-Za-z0-9_-]+", candidate):
        raise ValueError("Invalid candidate ID")
    path = safe_report_path(path)
    rubric = load_rubric()
    wb = load_workbook(path) if path.exists() else create_workbook(rubric)
    try:
        check_workbook(wb)
        if has_candidate(wb, candidate):
            print(f"{candidate}: already in v2 report; skipped")
            return False
        answers = read_answers(data_folder() / candidate)
        result, source, calculated, total = latest_valid(candidate, answers, rubric)
        append_rows(wb, candidate, answers, result, rubric, calculated, total)
        save_safely(wb, path)
        print(f"{candidate}: exported v2 {source.name}; total {total:.2f} / 40")
        return True
    finally:
        wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--export-only", nargs="+", required=True, metavar="ID")
    parser.add_argument("--report", type=Path, default=REPORT_PATH)
    args = parser.parse_args(argv)
    failed = False
    for candidate in args.export_only:
        try:
            export_candidate(candidate, args.report)
        except (OSError, ValueError, KeyError, TypeError, ArithmeticError):
            print(f"{candidate}: not exported; check v2 result validity, report path and Excel file access")
            failed = True
    return int(failed)


if __name__ == "__main__":
    raise SystemExit(main())
